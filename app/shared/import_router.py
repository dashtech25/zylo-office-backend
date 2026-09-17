"""Analyse générique de fichiers tabulaires uploadés (XLSX) vers un tableau
de cellules texte — miroir en lecture de `export_router.py`. Réutilisable par
tout import de données tabulaires de l'app, pas seulement Zylo Liquid (P1-7,
audit module Stations 2026-09-16 : la table de calibration d'une cuve ne
pouvait être modifiée que par import CSV, jamais XLSX). Le fichier n'est
jamais persisté ni interprété métier ici — seule la structure ligne/colonne
est extraite ; c'est à l'appelant (ex. import de calibration) de valider et
d'interpréter les valeurs."""

import io

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel

from app.core.security import get_current_user

import_router = APIRouter()


class ParsedTable(BaseModel):
    rows: list[list[str]]


@import_router.post("/xlsx", response_model=ParsedTable, dependencies=[Depends(get_current_user)])
async def parse_xlsx(file: UploadFile) -> ParsedTable:
    content = await file.read()
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=422, detail="Fichier XLSX invalide ou corrompu.")
    sheet = workbook.active
    if sheet is None:
        return ParsedTable(rows=[])
    rows = [[("" if cell is None else str(cell)) for cell in row] for row in sheet.iter_rows(values_only=True)]
    return ParsedTable(rows=rows)
