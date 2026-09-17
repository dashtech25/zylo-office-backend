"""Génération de documents métier — mission « bon de commande + aperçu/
partage ». Un seul document généré pour l'instant (bon de commande),
architecture pensée pour en accueillir d'autres plus tard sans les
anticiper ici. Aucune donnée inventée : le `PurchaseOrder` ne porte pas de
prix (pas de rapprochement tarifaire avec `PriceHistory` à la commande),
le document généré ne mentionne donc ni prix unitaire ni TVA/TTC — seules
les informations réellement connues (produits commandés, volumes, dates,
fournisseur, société) y figurent. Plus de « cuve de destination » depuis la
refonte 2026-09-17 (la cuve se choisit à la livraison, jamais à la
commande) — un bon de commande liste ses lignes produit×volume."""

import io
from datetime import datetime

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.modules.zylo_liquid.models import FuelProduct, PurchaseOrder, PurchaseOrderLine, Station, Supplier

_LINE_STATUS_LABEL = {"open": "Ouverte", "partially_received": "Partiellement reçue", "received": "Reçue"}


def _fmt_date(value: datetime | None) -> str:
    return value.strftime("%d/%m/%Y") if value else "Non renseignée"


def _fmt_volume(value: float) -> str:
    return f"{value:,.0f} L".replace(",", " ")


def _company_lines(station: Station) -> list[str]:
    lines = [station.name]
    if station.address:
        lines.append(station.address)
    contact = " · ".join(part for part in [station.phone, station.email] if part)
    if contact:
        lines.append(contact)
    if station.taxId:
        lines.append(f"SIRET/RCS : {station.taxId}")
    return lines


def _supplier_lines(supplier: Supplier) -> list[str]:
    lines = [supplier.name]
    if supplier.address:
        lines.append(supplier.address)
    contact = " · ".join(part for part in [supplier.contactName, supplier.contactPhone, supplier.contactEmail] if part)
    if contact:
        lines.append(contact)
    if supplier.taxId:
        lines.append(f"SIRET/RCS : {supplier.taxId}")
    return lines


def generate_purchase_order_pdf(
    purchase_order: PurchaseOrder, lines: list[tuple[PurchaseOrderLine, FuelProduct]], supplier: Supplier, station: Station,
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm, leftMargin=18 * mm, rightMargin=18 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("BonCommandeTitle", parent=styles["Title"], fontSize=20, spaceAfter=2 * mm)
    label_style = ParagraphStyle("Label", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, textColor=colors.HexColor("#555555"))
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=13)

    elements: list = [Paragraph("BON DE COMMANDE", title_style)]
    elements.append(Paragraph(f"Commande N° {purchase_order.orderReference}", body_style))
    elements.append(Paragraph(f"Date : {_fmt_date(purchase_order.orderedAt)}", body_style))
    elements.append(Spacer(1, 6 * mm))

    parties_table = Table(
        [
            [Paragraph("Acheteur", label_style), Paragraph("Fournisseur", label_style)],
            [
                Paragraph("<br/>".join(_company_lines(station)), body_style),
                Paragraph("<br/>".join(_supplier_lines(supplier)), body_style),
            ],
        ],
        colWidths=[85 * mm, 85 * mm],
    )
    parties_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, 0), 4)]))
    elements.append(parties_table)
    elements.append(Spacer(1, 8 * mm))

    order_rows = [["Désignation", "Volume commandé", "Livraison attendue", "Statut"]]
    for line, fuel_product in lines:
        order_rows.append([fuel_product.name, _fmt_volume(line.orderedVolumeLiters), _fmt_date(purchase_order.expectedAt), _LINE_STATUS_LABEL.get(line.status, line.status)])
    order_table = Table(order_rows, colWidths=[55 * mm, 40 * mm, 40 * mm, 35 * mm])
    order_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d0d0d0")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    elements.append(order_table)
    elements.append(Spacer(1, 10 * mm))
    elements.append(Paragraph(f"Statut global de la commande : {_LINE_STATUS_LABEL.get(purchase_order.status, purchase_order.status)}", body_style))
    elements.append(Spacer(1, 16 * mm))
    elements.append(Paragraph("Bon pour accord — signature :", label_style))
    elements.append(Spacer(1, 18 * mm))
    elements.append(Paragraph(
        "Ce document ne mentionne pas de prix : la tarification n'est pas rattachée aux commandes dans Zylo Liquid.",
        ParagraphStyle("Footnote", parent=styles["Normal"], fontSize=7.5, textColor=colors.HexColor("#888888")),
    ))

    doc.build(elements)
    return buffer.getvalue()


def generate_purchase_order_docx(
    purchase_order: PurchaseOrder, lines: list[tuple[PurchaseOrderLine, FuelProduct]], supplier: Supplier, station: Station,
) -> bytes:
    doc = DocxDocument()

    title = doc.add_heading("BON DE COMMANDE", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run(f"Commande N° {purchase_order.orderReference}\n").bold = True
    meta.add_run(f"Date : {_fmt_date(purchase_order.orderedAt)}")

    parties = doc.add_table(rows=2, cols=2)
    parties.style = "Table Grid"
    parties.cell(0, 0).paragraphs[0].add_run("Acheteur").bold = True
    parties.cell(0, 1).paragraphs[0].add_run("Fournisseur").bold = True
    parties.cell(1, 0).text = "\n".join(_company_lines(station))
    parties.cell(1, 1).text = "\n".join(_supplier_lines(supplier))

    doc.add_paragraph()

    order_table = doc.add_table(rows=1 + len(lines), cols=4)
    order_table.style = "Table Grid"
    headers = ["Désignation", "Volume commandé", "Livraison attendue", "Statut"]
    for i, header in enumerate(headers):
        cell_run = order_table.cell(0, i).paragraphs[0].add_run(header)
        cell_run.bold = True
    for row_index, (line, fuel_product) in enumerate(lines, start=1):
        values = [fuel_product.name, _fmt_volume(line.orderedVolumeLiters), _fmt_date(purchase_order.expectedAt), _LINE_STATUS_LABEL.get(line.status, line.status)]
        for i, value in enumerate(values):
            order_table.cell(row_index, i).text = value

    doc.add_paragraph()
    doc.add_paragraph(f"Statut global de la commande : {_LINE_STATUS_LABEL.get(purchase_order.status, purchase_order.status)}")
    doc.add_paragraph()
    doc.add_paragraph()
    signature = doc.add_paragraph("Bon pour accord — signature :")
    signature.runs[0].bold = True
    doc.add_paragraph()
    doc.add_paragraph()

    footnote = doc.add_paragraph("Ce document ne mentionne pas de prix : la tarification n'est pas rattachée aux commandes dans Zylo Liquid.")
    footnote.runs[0].font.size = Pt(7.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
