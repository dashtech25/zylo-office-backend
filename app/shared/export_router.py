"""Génération générique de fichiers d'export (XLSX/DOCX) à partir d'un
tableau déjà construit par l'appelant (en-têtes + lignes). Le frontend
construit déjà les données à exporter côté client (filtres, agrégations déjà
en mémoire, ex. la liste des stations) — ce routeur ne fait que les mettre
en forme dans le format demandé, sans toucher à aucune donnée métier ni
base de données. Réutilisable par n'importe quel écran de liste de l'app,
pas seulement Zylo Liquid (P1-3, audit module Stations 2026-09-16 :
l'export de la liste des stations était limité au CSV — le CSV reste
généré côté client, inchangé, seuls XLSX et DOCX manquaient)."""

import io
import re

from docx import Document as DocxDocument
from fastapi import APIRouter, Depends
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from app.core.security import get_current_user

export_router = APIRouter()

_UNSAFE_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9_.-]+")


class ExportTableRequest(BaseModel):
    filename: str = Field(min_length=1, max_length=100, description="Sans extension — ajoutée par l'endpoint selon le format.")
    headers: list[str] = Field(min_length=1)
    rows: list[list[str]] = Field(default_factory=list)


def _safe_filename(name: str, extension: str) -> str:
    base = _UNSAFE_FILENAME_CHARS.sub("_", name).strip("_") or "export"
    return f"{base}.{extension}"


@export_router.post("/xlsx", dependencies=[Depends(get_current_user)])
async def export_table_xlsx(data: ExportTableRequest) -> Response:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(data.headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in data.rows:
        sheet.append(row)

    for col_index, header in enumerate(data.headers, start=1):
        column_values = [header] + [str(row[col_index - 1]) for row in data.rows if col_index - 1 < len(row)]
        width = min(60, max(len(v) for v in column_values) + 2)
        sheet.column_dimensions[get_column_letter(col_index)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    filename = _safe_filename(data.filename, "xlsx")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@export_router.post("/docx", dependencies=[Depends(get_current_user)])
async def export_table_docx(data: ExportTableRequest) -> Response:
    # Même convention que document_generation.py (bons de commande) — table
    # "Table Grid", jamais un deuxième style de tableau DOCX réinventé.
    doc = DocxDocument()
    doc.add_heading(data.filename, level=1)
    table = doc.add_table(rows=1, cols=len(data.headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for i, header in enumerate(data.headers):
        header_cells[i].text = header
        for paragraph in header_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in data.rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if i < len(cells):
                cells[i].text = str(value)

    buffer = io.BytesIO()
    doc.save(buffer)
    filename = _safe_filename(data.filename, "docx")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
